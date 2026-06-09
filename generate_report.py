import sys
import subprocess
import os

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 1. Trang bìa
doc.add_paragraph('\n\n\n')
title = doc.add_paragraph('BÁO CÁO MÔN HỌC/DỰ ÁN')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(24)
    run.bold = True

doc.add_paragraph('\n\n')
proj = doc.add_paragraph('TÊN DỰ ÁN: ỨNG DỤNG ĐẶT LỊCH KHÁM SỨC KHỎE ONLINE TÍCH HỢP AI')
proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in proj.runs:
    run.font.size = Pt(18)
    run.bold = True

doc.add_paragraph('\n\n\n\n\n')
info = doc.add_paragraph('Sinh viên thực hiện: [NHẬP TÊN CỦA BẠN]\n\nMSSV: [NHẬP MÃ SỐ SINH VIÊN]')
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in info.runs:
    run.font.size = Pt(14)
    run.bold = True

doc.add_page_break()

# 2. Tóm tắt
doc.add_heading('Tóm tắt (Abstract)', level=1)
doc.add_paragraph('Dự án xây dựng hệ thống đặt lịch khám sức khỏe trực tuyến sử dụng ngôn ngữ lập trình Python và framework Streamlit. Hệ thống cung cấp giải pháp toàn diện bao gồm: tra cứu chuyên khoa dựa trên triệu chứng bằng trợ lý ảo AI (Google Gemini), kiểm tra trùng lịch thông minh, gợi ý thời gian thay thế, và tự động gửi email xác nhận/nhắc lịch hẹn thông qua SMTP. Ứng dụng giúp số hóa quy trình khám bệnh, tối ưu hóa thời gian đặt lịch, giảm tải cho bộ phận tổng đài và tiết kiệm thời gian cho cả người bệnh và cơ sở y tế.')

# 3. Mục lục
doc.add_heading('Mục lục', level=1)
doc.add_paragraph('1. Giới thiệu bài toán\n2. Kết quả đạt được\n3. Thảo luận, Ưu nhược điểm\n4. Kết luận, Hướng phát triển\n5. Tài liệu tham khảo\n6. Phụ lục (Mã nguồn)\n\n(Lưu ý: Người viết có thể sử dụng chức năng References -> Table of Contents của MS Word để tạo lại mục lục tự động và đánh số trang chính xác)')

doc.add_page_break()

# 4. Giới thiệu bài toán
doc.add_heading('1. Giới thiệu bài toán', level=1)
doc.add_paragraph('Ngày nay, nhu cầu chăm sóc sức khỏe của người dân ngày càng tăng cao. Tuy nhiên, quy trình đặt lịch khám truyền thống (gọi điện thoại qua tổng đài hoặc xếp hàng lấy số trực tiếp) thường tốn kém thời gian và dễ gây quá tải cục bộ. Hơn nữa, nhiều bệnh nhân khi có triệu chứng bệnh (ví dụ: đau ngực, mờ mắt, phát ban) lại không biết chính xác mình cần thăm khám ở chuyên khoa nào, dẫn đến việc mất thời gian chuyển tuyến khám.')
doc.add_paragraph('Để giải quyết vấn đề này, dự án "Booking Health Online" được phát triển nhằm cung cấp một nền tảng web trực quan, cho phép người dùng đặt lịch khám dễ dàng. Điểm nổi bật và khác biệt của hệ thống là khả năng ứng dụng Trí tuệ nhân tạo (Generative AI) để phân tích triệu chứng của người dùng theo ngôn ngữ tự nhiên và tự động đưa ra gợi ý chuyên khoa phù hợp.')

# 5. Kết quả
doc.add_heading('2. Kết quả đạt được', level=1)
doc.add_paragraph('Dự án đã phát triển và triển khai thành công ứng dụng web với các tính năng hoàn thiện sau:')
doc.add_paragraph('- Giao diện người dùng: Được xây dựng bằng Streamlit với thiết kế hiện đại, responsive, và thân thiện. Có hệ thống định vị bước (Step indicator) rõ ràng.')
doc.add_paragraph('- Trợ lý AI y tế: Tích hợp Google Gemini API để phân tích triệu chứng. Hệ thống có cơ chế dự phòng (fallback) chuyển sang chế độ Offline tự động bằng Rule-based nếu gặp sự cố giới hạn quota API.')
doc.add_paragraph('- Quản lý đặt lịch thông minh: Thuật toán kiểm tra xung đột thời gian (conflict resolution) theo thời gian thực. Tự động tính toán các khung giờ chưa diễn ra trong ngày hôm nay, và gợi ý khung giờ thay thế nếu bác sĩ đã kín lịch.')
doc.add_paragraph('- Thông báo tự động: Hệ thống tự động biên dịch thư nội dung và gửi email xác nhận ngay khi đặt lịch thành công, cùng module chạy nền hỗ trợ gửi email nhắc lịch hẹn cho bệnh nhân vào ngày mai.')
doc.add_paragraph('- Triển khai thực tế: Mã nguồn đã được cấu hình quản lý phiên bản với Git và triển khai thành công (deploy) lên nền tảng đám mây Streamlit Community Cloud.')

# 6. Thảo luận, Ưu nhược điểm
doc.add_heading('3. Thảo luận, Ưu nhược điểm', level=1)
doc.add_heading('3.1. Ưu điểm', level=2)
doc.add_paragraph('- Giải pháp Tích hợp AI thông minh giúp giải quyết tốt bài toán tư vấn phân luồng bệnh nhân mà các hệ thống truyền thống thiếu sót.')
doc.add_paragraph('- Luồng xử lý thời gian (logic nghiệp vụ) chặt chẽ, không cho phép đặt các giờ đã qua ở hiện tại, không gây lỗi trùng lịch.')
doc.add_paragraph('- Giao diện trực quan, trải nghiệm người dùng tối ưu, chi phí triển khai hệ thống thấp (mô hình serverless web).')

doc.add_heading('3.2. Nhược điểm', level=2)
doc.add_paragraph('- Hệ thống lưu trữ hiện tại đang sử dụng cấu trúc file tĩnh (CSV/JSON) để tiện cho việc demo, chưa đáp ứng được tính toàn vẹn (ACID) cho quy mô dữ liệu lớn và truy cập đồng thời.')
doc.add_paragraph('- Tốc độ phản hồi ở khâu phân tích triệu chứng phụ thuộc vào đường truyền mạng và giới hạn của API Google Gemini.')

# 7. Kết luận, Hướng phát triển
doc.add_heading('4. Kết luận, Hướng phát triển', level=1)
doc.add_heading('4.1. Kết luận', level=2)
doc.add_paragraph('Dự án đã hoàn thành toàn bộ mục tiêu đề ra, xây dựng thành công một hệ thống đặt lịch khám bệnh trực tuyến thông minh. Dự án chứng minh tính khả thi của việc tích hợp LLM (Large Language Models) vào các dịch vụ y tế quy mô vừa và nhỏ để nâng cao trải nghiệm người bệnh.')

doc.add_heading('4.2. Hướng phát triển', level=2)
doc.add_paragraph('- Nâng cấp Cơ sở dữ liệu: Chuyển đổi từ file CSV sang Hệ quản trị cơ sở dữ liệu quan hệ (PostgreSQL/MySQL) hoặc NoSQL (MongoDB).')
doc.add_paragraph('- Tích hợp Thanh toán: Kết nối các cổng thanh toán trực tuyến (VNPay, MoMo) để người dùng thanh toán/đặt cọc phí khám bệnh ngay trên web.')
doc.add_paragraph('- Quản lý Admin: Xây dựng Dashboard riêng biệt nâng cao hơn cho bác sĩ và quản trị viên xem thống kê doanh thu.')

# 8. Tài liệu tham khảo
doc.add_heading('5. Tài liệu tham khảo', level=1)
doc.add_paragraph('[1]. Tài liệu chính thức thư viện Streamlit (https://docs.streamlit.io/)')
doc.add_paragraph('[2]. Tài liệu Google Generative AI API (Gemini)')
doc.add_paragraph('[3]. Thư viện chuẩn Python: smtplib, email.mime')

# 9. Phụ lục (mã nguồn)
doc.add_page_break()
doc.add_heading('6. Phụ lục (Trích xuất Mã nguồn)', level=1)
doc.add_paragraph('Dưới đây là một số đoạn mã nguồn cốt lõi thể hiện logic quan trọng của hệ thống:')

doc.add_heading('6.1. Xử lý khung giờ trống theo thời gian thực (app.py)', level=2)
code_1 = doc.add_paragraph('''
# Lọc khung giờ để không hiển thị giờ đã qua trong ngày hiện tại
while current <= end_t:
    slot_str = current.strftime("%H:%M")
    
    if appt_date == date.today():
        slot_time_today = current.replace(year=now.year, month=now.month, day=now.day)
        if slot_time_today > now:
            time_slots.append(slot_str)
    else:
        time_slots.append(slot_str)
        
    current += timedelta(minutes=TIME_SLOT_DURATION_MINUTES)
''')
code_1.style = 'No Spacing'
for run in code_1.runs:
    run.font.name = 'Courier New'

doc.add_heading('6.2. Tích hợp AI tư vấn triệu chứng (ai_service.py)', level=2)
code_2 = doc.add_paragraph('''
def get_gemini_response(prompt: str) -> str:
    # Gửi triệu chứng bệnh nhân đến mô hình Gemini để phân tích
    client = genai.Client(api_key=GEMINI_API_KEY)
    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=[system_instruction, prompt]
    )
    return response.text
''')
code_2.style = 'No Spacing'
for run in code_2.runs:
    run.font.name = 'Courier New'

doc.save('Bao_cao_Du_an_Booking_Health.docx')
print("Successfully generated Bao_cao_Du_an_Booking_Health.docx")
